from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "templates" / "resume-template.docx"

def font(run, size=10, bold=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold

def line(doc, text, size=10, bold=False, centered=False, after=2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    font(r, size, bold)
    return p

def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper())
    font(r, 10, True)

def entry(doc, title, detail, bullets):
    p = line(doc, "", after=1)
    r = p.add_run(title)
    font(r, 10, True)
    r = p.add_run(" | " + detail)
    font(r, 10)
    for bullet in bullets:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(bullet)
        font(r, 10)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.45)
section.bottom_margin = Inches(0.45)
section.left_margin = Inches(0.58)
section.right_margin = Inches(0.58)
styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(10)

line(doc, "Your Name", size=16, bold=True, centered=True, after=1)
line(doc, "City, State | email@example.com | (000) 000-0000 | linkedin.com/in/yourname", centered=True, after=5)
heading(doc, "Education")
entry(doc, "University Name", "City, State | Expected Month Year", ["Degree, major or concentration | GPA if you choose to include it", "Relevant coursework, honors, or activities if useful for the target role"])
heading(doc, "Experience")
entry(doc, "Organization", "Role | Month Year - Month Year", ["Start each bullet with a verb. State what you did, who or what it supported, and a verified result when available.", "Use concrete scope only when you can support it. Replace this guidance with your own evidence."])
heading(doc, "Projects")
entry(doc, "Project name", "Tools or focus area | Month Year", ["Explain the problem, your contribution, and the evidence of what changed or what you learned."])
heading(doc, "Leadership and activities")
entry(doc, "Organization", "Position | Month Year - Present", ["Use this section for sustained responsibility, community work, or leadership that adds a distinct strength."])
heading(doc, "Skills")
line(doc, "Technical: tools you can use honestly | Languages: languages and proficiency | Other: relevant methods or certifications", after=0)

doc.core_properties.title = "Internship Resume Template"
doc.core_properties.author = "Internship Search Starter"
doc.core_properties.subject = "Flexible one page resume template"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
